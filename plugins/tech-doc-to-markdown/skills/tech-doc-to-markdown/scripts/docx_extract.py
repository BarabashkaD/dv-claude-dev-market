"""Draft Markdown from a .docx in document order: headings, paragraphs, list items, tables and image placeholders.

Usage:
  python docx_extract.py INPUT.docx --out DRAFT.md [--assets DIR]

Requires python-docx. For .doc/.rtf/.odt convert to .docx first (office_convert.ps1 or LibreOffice).
The draft is a starting point. python-docx does not see text boxes, SmartArt, drawing-canvas shapes or
equations; take those from rendered pages (export a PDF and run survey.py). Embedded pictures are exported to
--assets for inspection; diagrams among them still get transcribed as text.
"""
import argparse
import pathlib
import re

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def run_text(paragraph):
    out = []
    for run in paragraph.runs:
        t = run.text
        if not t:
            continue
        if run.bold and t.strip():
            t = f"**{t}**"
        elif run.italic and t.strip():
            t = f"*{t}*"
        out.append(t)
    return "".join(out).replace("****", "")


def cell_text(cell):
    return "<br>".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).replace("|", "\\|")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("--out", required=True)
    ap.add_argument("--assets")
    args = ap.parse_args()

    document = docx.Document(args.input)
    assets = pathlib.Path(args.assets) if args.assets else None
    if assets:
        assets.mkdir(parents=True, exist_ok=True)
    md, image_no, tables = [], 0, 0

    def blank():
        if md and md[-1] != "":
            md.append("")

    for block in blocks(document):
        if isinstance(block, Table):
            tables += 1
            rows = [[cell_text(c) for c in row.cells] for row in block.rows]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            blank()
            md.append(f"<!-- table {tables} -->")
            md.append("| " + " | ".join(rows[0]) + " |")
            md.append("|" + "---|" * width)
            md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            md.append("")
            continue

        style = block.style.name if block.style is not None else ""
        for rid in block._p.xpath(".//a:blip/@r:embed"):
            image_no += 1
            name = f"image-{image_no:03d}"
            try:
                part = document.part.related_parts[rid]
                ext = pathlib.Path(str(part.partname)).suffix or ".bin"
                if assets:
                    (assets / f"{name}{ext}").write_bytes(part.blob)
                md.append(f"<!-- {name}{ext}: embedded picture; transcribe as text unless it is a photo -->")
            except KeyError:
                md.append(f"<!-- {name}: linked picture {rid} not embedded -->")
        text = run_text(block).strip()
        if not text:
            continue
        m = re.match(r"^(?:Heading|Заголовок)\s*(\d)", style)
        is_list = style.startswith("List")
        if not is_list and md and re.match(r"^(- |1\. )", md[-1]):
            blank()
        if m:
            blank()
            md += [f"{'#' * min(int(m.group(1)) + 1, 6)} {block.text.strip()}", ""]
        elif style == "Title":
            blank()
            md += [f"# {block.text.strip()}", ""]
        elif style.startswith("List Number"):
            md.append(f"1. {text}")
        elif is_list:
            md.append(f"- {text}")
        else:
            md += [text, ""]

    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")
    print(f"written {out}: {len(md)} lines, {tables} tables, {image_no} pictures"
          + (f" exported to {assets}" if assets and image_no else ""))


if __name__ == "__main__":
    main()
